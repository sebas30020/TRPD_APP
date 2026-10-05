"""Presentación PowerPoint con las señales temporales de todas las mediciones.

Una diapositiva por disparo (segmento) con 4 gráficos apilados y eje de tiempo
compartido: Impulso (CH1), HFCT (CH2), Antena 1 – cercana (CH3) y Antena 2 –
lejana (CH4). Señales igual que en la app: CH2..CH4 filtradas según
canales.chX.filtro de metadata.yaml (por defecto HP 5 MHz / HP 200 MHz), CH1 sin
filtrar, ventana -5..30 µs. Cada título indica nº de vacuolas, diámetros y tensión,
leídos de la jerarquía <probeta>/<X>kV (p. ej. 3v_2mm3mm4mm_1/13kV).
Con --crudo se dibujan las señales tal como las registró el osciloscopio, sin
ningún filtro digital (salida por defecto senales_med_proced_confuse_crudo.pptx).

Uso:
    python generar_presentacion.py [<raiz> ...] [--salida <ruta.pptx>]
                                   [--dpi 150] [--solo <subcadena de la ruta>] [--crudo]
Cada <raiz> puede ser una carpeta de probetas (<raiz>/<probeta>/<X>kV) o una
probeta (<raiz>/<X>kV). Orden: por nº de vacuolas y diámetros, luego por tensión
creciente (así un set posterior medido a menor tensión queda antes).
Dependencias: requirements_reportes.txt (python-pptx, matplotlib).
"""

import argparse
import datetime
import io
import os
import sys

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, AQUI)

import app  # noqa: E402  (lectura filtrada, caché y diezmado de la app)
import filtros  # noqa: E402
import rutas  # noqa: E402
import tema  # noqa: E402
from generate_metadata import inferir_parametros  # noqa: E402

PROYECTO = os.path.dirname(AQUI)
RAICES_DEFECTO = [
    os.path.join(PROYECTO, "mediciones", "med_proced_confuse"),
    os.path.join(PROYECTO, "mediciones", "med_proced", "1v_2mm_01"),
]
SALIDA_DEFECTO = os.path.join(PROYECTO, "reportes", "senales_med_proced_confuse.pptx")
SALIDA_CRUDO = os.path.join(PROYECTO, "reportes", "senales_med_proced_confuse_crudo.pptx")

CANALES = [
    ("ch1", "Impulso\nCH1", "V"),
    ("ch2", "HFCT\nCH2", "mV"),
    ("ch3", "Antena 1\n(cercana)\nCH3", "mV"),
    ("ch4", "Antena 2\n(lejana)\nCH4", "mV"),
]
PUNTOS_POR_TRAZA = 6000          # diezmado min-max: conserva todos los picos
ANCHO, ALTO = Inches(13.333), Inches(7.5)
GRIS = RGBColor(0x5F, 0x6B, 0x7A)
TINTA = RGBColor(0x1F, 0x29, 0x33)


# ---------------- Mediciones ----------------

def _carpetas_medicion(raiz):
    """Carpetas con chN.h5 en `raiz` o hasta dos niveles por debajo."""
    if app.canales_presentes(raiz):
        return [raiz]
    encontradas = []
    for sub in os.listdir(raiz):
        d = os.path.join(raiz, sub)
        if os.path.isdir(d):
            encontradas += [d] if app.canales_presentes(d) else [
                os.path.join(d, x) for x in os.listdir(d)
                if os.path.isdir(os.path.join(d, x)) and app.canales_presentes(os.path.join(d, x))]
    return encontradas


def descubrir(raices, solo=None):
    """Mediciones bajo `raices`, ordenadas por (nº de vacuolas, diámetros), luego
    por tensión creciente y por último por nombre de probeta."""
    meds = []
    for raiz in raices:
        for d in _carpetas_medicion(raiz):
            inf = inferir_parametros(d)
            probeta = os.path.basename(os.path.dirname(d))
            meds.append((probeta, inf.get("tension_sec_kv") or 0.0, d, inf, raiz))
    meds.sort(key=lambda m: (m[3].get("nro_vacuolas") or 99, tuple(m[3].get("diametros") or ()),
                             m[1], rutas.orden_natural(m[0])))
    if solo:
        clave = solo.replace("/", "\\").lower()
        meds = [m for m in meds if clave in m[2].replace("/", "\\").lower()]
    return meds


def _diametros(lista):
    nums = [f"{float(d):g}" for d in (lista or [])]
    if not nums:
        return "Ø ?"
    if len(nums) == 1:
        return f"Ø {nums[0]} mm"
    return f"Ø {', '.join(nums[:-1])} y {nums[-1]} mm"


def describir(carpeta, inf, raiz):
    n = inf.get("nro_vacuolas")
    vac = f"{n} vacuola" + ("" if n == 1 else "s") if n else (inf.get("codigo_probeta") or "?")
    kv = inf.get("tension_sec_kv")
    meta = app.obtener_metadata(carpeta)
    return {
        "titulo": f"{vac} · {_diametros(inf.get('diametros'))} · " + (f"{kv:g} kV" if kv else "? kV"),
        "vacuolas": vac,
        "diametros": _diametros(inf.get("diametros")),
        "tension": f"{kv:g} kV" if kv else "?",
        "probeta": inf.get("codigo_probeta") or os.path.basename(os.path.dirname(carpeta)),
        # set tal como está escrito en la carpeta ('01' no se convierte en 1)
        "set": (inf.get("codigo_probeta") or "").rsplit("_", 1)[-1] if inf.get("codigo_probeta") else None,
        "ruta": os.path.relpath(carpeta, os.path.dirname(raiz)),
        "fecha": (meta.get("experimento") or {}).get("fecha_hora") or "",
        "nsegs": app.n_segmentos(carpeta),
        "filtros": " · ".join(f"{ch.upper()} {filtros.etiqueta(app.spec_filtro(carpeta, ch))}"
                              for ch in ("ch2", "ch3", "ch4") if ch in app.canales_presentes(carpeta)),
    }


# ---------------- Figuras ----------------

def _senal(carpeta, ch, seg, filtrado=True):
    t, v = app.cargar_segmento(carpeta, ch, seg, filtrado=filtrado)   # CH1 nunca se filtra
    return t, (v / 1000.0 if ch == "ch1" else v)                  # CH1 en V


def limites_y(carpeta, filtrado=True):
    """(ymin, ymax) por canal, comunes a todos los disparos de la medición."""
    lim = {}
    for ch, _, _ in CANALES:
        if ch not in app.canales_presentes(carpeta):
            continue
        lo, hi = np.inf, -np.inf
        for s in range(1, app.n_segmentos(carpeta) + 1):
            _, v = _senal(carpeta, ch, s, filtrado)
            if v.size:
                lo, hi = min(lo, float(v.min())), max(hi, float(v.max()))
        if np.isfinite(lo):
            m = 0.05 * (hi - lo or 1.0)
            lim[ch] = (lo - m, hi + m)
    return lim


def figura_segmento(carpeta, seg, lim, dpi, filtrado=True):
    fig, ejes = plt.subplots(4, 1, sharex=True, figsize=(13.33, 6.3), dpi=dpi)
    for ax, (ch, nombre, unidad) in zip(ejes, CANALES):
        if ch in app.canales_presentes(carpeta):
            t, v = _senal(carpeta, ch, seg, filtrado)
            t, v = app.decimar_minmax(t, v, PUNTOS_POR_TRAZA)
            ax.plot(t, v, color=tema.COLORES_CANALES.get(ch, "#2f5d8a"), lw=0.7)
            if ch in lim:
                ax.set_ylim(*lim[ch])
        else:
            ax.text(0.5, 0.5, "canal no disponible", transform=ax.transAxes,
                    ha="center", va="center", color="#9aa5b1")
        ax.set_ylabel(f"{nombre} [{unidad}]", fontsize=8.5, linespacing=1.1)
        ax.grid(True, color="#e5e7eb", lw=0.6)
        ax.tick_params(labelsize=8)
        for lado in ("top", "right"):
            ax.spines[lado].set_visible(False)
    ejes[-1].set_xlim(app.T_MIN, app.T_MAX)
    ejes[-1].set_xlabel("Tiempo [µs]", fontsize=9)
    fig.align_ylabels(ejes)
    fig.tight_layout(h_pad=0.4)
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=dpi)
    plt.close(fig)
    buf.seek(0)
    return buf


# ---------------- Diapositivas ----------------

def _texto(slide, x, y, w, h, texto, tam, negrita=False, color=TINTA):
    caja = slide.shapes.add_textbox(x, y, w, h)
    tf = caja.text_frame
    tf.word_wrap = True
    for i, linea in enumerate(texto.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        r = p.add_run()
        r.text = linea
        r.font.size, r.font.bold, r.font.color.rgb = Pt(tam), negrita, color
    return caja


def _nueva(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])   # en blanco


def portada(prs, meds, raices, total, filtrado=True):
    s = _nueva(prs)
    _texto(s, Inches(0.8), Inches(1.6), Inches(11.7), Inches(1.0),
           "Señales temporales de las mediciones" + ("" if filtrado else " — sin filtrar"), 36, True)
    _texto(s, Inches(0.8), Inches(2.55), Inches(11.7), Inches(0.6),
           "Fuentes: " + " · ".join(os.path.relpath(r, os.path.join(PROYECTO, "mediciones")) for r in raices),
           14, color=GRIS)
    _texto(s, Inches(0.8), Inches(3.3), Inches(11.7), Inches(2.5),
           f"{len(meds)} mediciones · {total} disparos · una diapositiva por disparo\n"
           "Canales: Impulso (CH1) · HFCT (CH2) · Antena 1 – cercana (CH3) · Antena 2 – lejana (CH4)\n"
           + ("CH2..CH4 con los filtros digitales de la app (Butterworth orden 4, fase cero); CH1 sin filtrar\n"
              if filtrado else "Señales crudas, tal como las registró el osciloscopio: ningún canal filtrado\n")
           + f"Ventana {app.T_MIN:g} a {app.T_MAX:g} µs · escala vertical común a los disparos de cada medición\n"
           f"Generado el {datetime.date.today().isoformat()} con TRPD_APP/generar_presentacion.py",
           16, color=GRIS)


def indice(prs, descs):
    s = _nueva(prs)
    _texto(s, Inches(0.6), Inches(0.35), Inches(12), Inches(0.7), "Índice de mediciones", 26, True)
    cols = ["Probeta", "Vacuolas", "Diámetros", "Tensión", "Disparos", "Fecha"]
    alto_fila = min(Inches(0.38), int(Inches(5.9) / (len(descs) + 1)))
    tabla = s.shapes.add_table(len(descs) + 1, len(cols), Inches(0.6), Inches(1.2),
                               Inches(12.1), alto_fila * (len(descs) + 1)).table
    for fila in tabla.rows:
        fila.height = alto_fila
    for j, c in enumerate(cols):
        tabla.cell(0, j).text = c
    for i, d in enumerate(descs, start=1):
        valores = [d["probeta"], d["vacuolas"], d["diametros"], d["tension"], str(d["nsegs"]), d["fecha"]]
        for j, val in enumerate(valores):
            tabla.cell(i, j).text = val
    for fila in tabla.rows:
        for celda in fila.cells:
            for p in celda.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(12 if len(descs) <= 13 else 10)


def separador(prs, d):
    s = _nueva(prs)
    _texto(s, Inches(0.8), Inches(2.4), Inches(11.7), Inches(1.2), d["titulo"], 34, True)
    set_txt = f" (set {d['set']})" if d["set"] is not None else ""
    _texto(s, Inches(0.8), Inches(3.6), Inches(11.7), Inches(2.0),
           f"Probeta {d['probeta']}{set_txt}\n{d['ruta']}\n"
           f"{d['nsegs']} disparos" + (f" · adquirido el {d['fecha']}" if d["fecha"] else ""),
           18, color=GRIS)


def diapositiva_segmento(prs, d, seg, png, filtrado=True):
    s = _nueva(prs)
    _texto(s, Inches(0.4), Inches(0.12), Inches(12.5), Inches(0.55),
           f"{d['titulo']} — Disparo {seg} / {d['nsegs']}", 22, True)
    set_txt = f" (set {d['set']})" if d["set"] is not None else ""
    _texto(s, Inches(0.4), Inches(0.6), Inches(12.5), Inches(0.35),
           f"Probeta {d['probeta']}{set_txt} · {d['ruta']}", 11, color=GRIS)
    s.shapes.add_picture(png, Inches(0.0), Inches(0.92), width=ANCHO)
    _texto(s, Inches(0.4), Inches(7.12), Inches(12.5), Inches(0.3),
           (f"{d['filtros']} (Butterworth orden 4, fase cero) · CH1 sin filtrar · " if filtrado
            else "Señales crudas del osciloscopio, sin filtros · ")
           + f"Ventana {app.T_MIN:g}…{app.T_MAX:g} µs", 9, color=GRIS)


# ---------------- Principal ----------------

def generar(raices, salida, dpi=150, solo=None, filtrado=True):
    meds = descubrir(raices, solo)
    if not meds:
        raise SystemExit(f"No hay mediciones en {raices}" + (f" que contengan '{solo}'" if solo else ""))
    descs = [describir(m[2], m[3], os.path.dirname(m[4]) if m[4].endswith(m[0]) else m[4])
             for m in meds]
    total = sum(d["nsegs"] for d in descs)
    prs = Presentation()
    prs.slide_width, prs.slide_height = ANCHO, ALTO
    portada(prs, meds, raices, total, filtrado)
    indice(prs, descs)
    for i, ((_, _, carpeta, _, _), d) in enumerate(zip(meds, descs), start=1):
        separador(prs, d)
        lim = limites_y(carpeta, filtrado)
        for seg in range(1, d["nsegs"] + 1):
            print(f"medición {i}/{len(meds)} ({d['ruta']}), disparo {seg}/{d['nsegs']}", flush=True)
            diapositiva_segmento(prs, d, seg, figura_segmento(carpeta, seg, lim, dpi, filtrado), filtrado)
    os.makedirs(os.path.dirname(os.path.abspath(salida)), exist_ok=True)
    tmp = salida + ".tmp"
    prs.save(tmp)
    os.replace(tmp, salida)   # sin archivos a medias si se interrumpe
    print(f"Presentación guardada: {salida} ({len(prs.slides)} diapositivas, "
          f"{os.path.getsize(salida) / 1e6:.1f} MB)")
    return salida


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("raices", nargs="*", default=RAICES_DEFECTO)
    ap.add_argument("--salida", default=None)
    ap.add_argument("--dpi", type=int, default=150)
    ap.add_argument("--solo", default=None, help="genera solo las mediciones cuya ruta contenga este texto")
    ap.add_argument("--crudo", action="store_true", help="señales sin ningún filtro digital")
    a = ap.parse_args()
    salida = a.salida or (SALIDA_CRUDO if a.crudo else SALIDA_DEFECTO)
    generar(a.raices, salida, a.dpi, a.solo, filtrado=not a.crudo)


if __name__ == "__main__":
    main()
