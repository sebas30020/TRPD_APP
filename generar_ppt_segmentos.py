"""Presentación (.pptx) con segmentos (disparos) elegidos de las mediciones del reporte 4.

Una diapositiva por segmento: a la izquierda el segmento completo con todos los canales
(CH1 impulso y los sensores CH2..CH4), a la derecha un zoom a la zona de las descargas.
Los peaks marcados son los del reporte 4 (trigger guardado + ediciones de metadata.yaml,
t_abs = t − t10 del disparo, sin calibración de retardo; V_p = máximo en la marca).

Uso:
    python generar_ppt_segmentos.py [--salida reportes/reporte4/segmentos_seleccionados.pptx]
Los segmentos se editan en SELECCION.
"""

import argparse
import io
import os
import sys

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, AQUI)

import app  # noqa: E402
import generar_reporte4 as g4  # noqa: E402
import reproducir_trpd  # noqa: E402
import tema  # noqa: E402

PROBETA = "3v_4mm_1"
# kV -> segmentos, en el orden elegido
SELECCION = {
    11: [10, 14, 15],
    12: [9],
    13: [],
    14: [2],
    15: [11, 9, 10],
    16: [2, 8, 15, 9, 11],
    17: [1, 7, 9, 12, 14, 15],
    18: [2, 1, 6, 4, 9, 11],
    19: [2, 6, 8],
    20: [1, 6, 10, 12, 13],
}
SALIDA_DEFECTO = os.path.join(g4.SALIDA_DEFECTO, "segmentos_seleccionados.pptx")
CANALES = ("ch1", "ch2", "ch3", "ch4")
MARGEN_ZOOM = (0.08, 0.15)   # µs antes del primer peak y después del último


def figura(carpeta, seg, res, t10):
    """Figura 4x2: segmento completo (izq.) y zoom a las descargas (der.)."""
    peaks = {}
    for c, r in res.items():
        m = np.asarray(r["cap"]["seg"]) == seg
        peaks[c] = (np.asarray(r["t_abs"])[m], np.asarray(r["cap"]["v_peak"], dtype=float)[m])
    t_todos = np.concatenate([p[0] for p in peaks.values()]) if peaks else np.array([])
    if t_todos.size:
        z0, z1 = float(t_todos.min()) - MARGEN_ZOOM[0], float(t_todos.max()) + MARGEN_ZOOM[1]
    else:   # sin peaks: zoom en torno al máximo |v| de los sensores tras t_abs = 0
        mejor, tc = -1.0, 1.0
        for c in res:
            t, v = app.cargar_segmento(carpeta, c, seg, filtrado=True)
            t = np.asarray(t) - t10
            sel = t > 0.0
            rel = np.abs(np.asarray(v)[sel]) / max(float(np.max(np.abs(v))), 1e-12)
            i = int(np.argmax(np.abs(np.asarray(v)[sel])))
            if rel[i] >= mejor:
                mejor, tc = rel[i], float(t[sel][i])
        z0, z1 = tc - 0.25, tc + 0.25
    x1 = max(4.0, float(np.ceil(z1 + 0.5)))

    fig, ejes = plt.subplots(4, 2, figsize=(13.0, 6.1), sharex="col",
                             gridspec_kw={"width_ratios": [1.35, 1.0], "hspace": 0.12, "wspace": 0.16})
    for fila, c in enumerate(CANALES):
        if c not in app.canales_presentes(carpeta):
            continue
        t, v = app.cargar_segmento(carpeta, c, seg, filtrado=(c != "ch1"))
        t = np.asarray(t) - t10
        color = tema.COLORES_CANALES[c]
        if c == "ch1":
            v = np.asarray(v) / max(float(np.max(np.abs(v))), 1e-12)
            nombre, unidad = "CH1 Impulso", "norm."
        else:
            r = res.get(c)
            nombre = f"{c.upper()} {r['fila']['sensor']}" if r else c.upper()
            unidad = "mV"
        for k, (ax, (a, b)) in enumerate(zip(ejes[fila], ((-0.5, x1), (z0, z1)))):
            sel = (t >= a) & (t <= b)
            ax.plot(t[sel], np.asarray(v)[sel], color=color, lw=0.6 if k == 0 else 0.8)
            if c in peaks and peaks[c][0].size:
                tp, vp = peaks[c]
                dentro = (tp >= a) & (tp <= b)
                ax.plot(tp[dentro], vp[dentro], "x", color="#c0392b", ms=7 if k else 5, mew=1.6)
                if k == 1:
                    for ti, vi in zip(tp[dentro], vp[dentro]):
                        ax.annotate(f"{ti:.3f}", (ti, vi), textcoords="offset points", xytext=(3, 3),
                                    fontsize=6.5, color="#c0392b")
            ax.set_xlim(a, b)
            ax.axvline(0.0, color=tema.LINEA_T10, lw=0.7, ls=":")
            ax.grid(True, color="#e5e7eb", lw=0.5)
            ax.tick_params(labelsize=7)
            for lado in ("top", "right"):
                ax.spines[lado].set_visible(False)
            if k == 0:
                ax.set_ylabel(f"{nombre}\n[{unidad}]", fontsize=7.5)
                ax.axvspan(z0, z1, color="#f4d03f", alpha=0.18, lw=0)
    ejes[0, 0].set_title("Segmento completo", fontsize=9, loc="left")
    ejes[0, 1].set_title(f"Zoom a las descargas ({z0:.2f} – {z1:.2f} µs)", fontsize=9, loc="left")
    for ax in ejes[-1]:
        ax.set_xlabel("t_abs [µs] (t10 del impulso = 0)", fontsize=8)
    fig.subplots_adjust(left=0.08, right=0.99, top=0.95, bottom=0.08)
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=170)
    plt.close(fig)
    buf.seek(0)
    return buf, peaks


def texto(slide, x, y, w, h, s, size, bold=False, color=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = s
    p.font.size = Pt(size)
    p.font.bold = bold
    if color is not None:
        p.font.color.rgb = color
    return tb


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--salida", default=SALIDA_DEFECTO)
    a = ap.parse_args()

    from pptx.dml.color import RGBColor
    gris = RGBColor(0x55, 0x55, 0x55)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    vacio = prs.slide_layouts[6]

    s = prs.slides.add_slide(vacio)
    texto(s, 0.8, 2.6, 11.5, 1.0, f"Segmentos seleccionados · {PROBETA}", 36, bold=True)
    texto(s, 0.8, 3.6, 11.5, 1.6,
          "Un segmento por diapositiva: todos los canales y zoom a la zona de descargas.\n"
          "× = peaks del reporte 4 (trigger guardado + ediciones manuales); t_abs = t − t10 del disparo.",
          16, color=gris)

    omitidos = []
    for kv, segs in SELECCION.items():
        carpeta = os.path.join(g4.MED, PROBETA, f"{kv}kV")
        if not segs:
            continue
        res = reproducir_trpd.resultado_medicion(carpeta, calibrado=False)
        nsegs = app.n_segmentos(carpeta)
        excl = app.segmentos_excluidos(carpeta)
        t10s = app.t10_por_segmento(carpeta)
        for seg in segs:
            if not 1 <= seg <= nsegs:
                omitidos.append(f"{kv} kV segmento {seg} (la medición tiene {nsegs})")
                continue
            img, peaks = figura(carpeta, seg, res, float(t10s[seg - 1]) if len(t10s) >= seg else 0.0)
            s = prs.slides.add_slide(vacio)
            tit = f"{kv} kV · segmento {seg}" + ("  (excluido del análisis)" if seg in excl else "")
            texto(s, 0.4, 0.15, 9.0, 0.6, tit, 26, bold=True)
            cuentas = " · ".join(f"{res[c]['fila']['sensor']}: {peaks[c][0].size}" for c in res if c in peaks)
            texto(s, 0.4, 0.75, 12.5, 0.4, f"N_PD por sensor — {cuentas}", 13, color=gris)
            s.shapes.add_picture(img, Inches(0.15), Inches(1.15), width=Inches(13.0))
            notas = [f"{PROBETA} {kv} kV, segmento {seg}. Peaks (t_abs µs / V_p V):"]
            for c in res:
                tp, vp = peaks.get(c, ([], []))
                notas.append(f"{res[c]['fila']['sensor']}: " + ", ".join(
                    f"{t:.3f} / {v / 1000:.3f}" for t, v in zip(tp, vp)))
            s.notes_slide.notes_text_frame.text = "\n".join(notas)
            print(f"{kv} kV segmento {seg}: {cuentas}")
    os.makedirs(os.path.dirname(a.salida), exist_ok=True)
    prs.save(a.salida)
    print(f"Presentación: {a.salida}")
    for o in omitidos:
        print(f"Omitido: {o}")


if __name__ == "__main__":
    main()
